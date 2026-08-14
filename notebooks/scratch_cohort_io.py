"""Subject-agnostic loading for the six-subject cohort (session S07).

The session 4/5 code is Rocky-shaped: two arrays named Anterior and Posterior,
one CMP pair, `.ns5` for broadband, one implant. None of that survives contact
with the cohort — Fisk's broadband is `.ns6` and its LFP `.ns3`, Rocky has two
implants that reuse the same two array labels, and Picasso has no Blackrock at
all. This module moves every one of those facts into `configs/subjects/*.json`
and reads them at run time.

It also validates probe geometry rather than trusting it. CLAUDE.md's standing
warning is that channel-order mismatch is silent and ruinous; the same is true
of *position* mismatch, which is subtler because the channel order can be
perfectly correct while the electrode sits in the wrong place on the grid.

**Which four cells of the 10x10 are unpopulated is a property of the individual
array, not of the array type.** Blackrock builds most arrays with the four
symmetric corners empty, but when shanks break during manufacture they rewire
surviving shanks from elsewhere to reach 96 channels, and the vacant cells move
accordingly. `SN 1025-004377` is such an array: `elec18` sits at top-left and
`elec8` at bottom-right, leaving `(8,9)` and `(9,8)` empty instead of `(0,9)`
and `(9,0)`. A validator that assumes the symmetric layout reports that as a
defect, and "repairing" it would move two electrodes to cells that hold no
electrode at all.

So geometry is checked against the array's own factory record — the pad-side
location grid in the `.xlsm` — and never against a canonical layout or against
a sibling array.

Run from repo root:

    uv run python notebooks/scratch_cohort_io.py --validate
    uv run python notebooks/scratch_cohort_io.py --regress

See:
- docs/cohort_plan.md
- docs/notes/utah_channel_mapping.md
"""

from __future__ import annotations

import argparse
import io
import json
import re
import struct
import warnings
import zipfile
import zlib
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

warnings.filterwarnings("ignore")

REPO = Path(__file__).resolve().parent.parent
CONFIG_DIR = REPO / "configs" / "subjects"
PROBE_DIR = REPO / "configs" / "probes"
ROCKY_PREIMPLANT = Path(r"D:\Claude Code\Rocky\preimplant")

# A Utah-96 sits on a 10x10 grid with 96 of the 100 cells populated. Which four
# are empty varies per array -- see the module docstring -- so it is reported,
# never assumed.
GRID = 10
N_ELECTRODES = 96
TYPICAL_VACANT = {(0, 0), (0, GRID - 1), (GRID - 1, 0), (GRID - 1, GRID - 1)}

# Pad-side location grid inside the factory workbook: columns AR..BA, rows
# 15..24. Printed top row is the highest cmp row, because cmp rows count up
# from the bottom.
PAD_SHEET = "Array Map with Automated Tester"
PAD_COL0, PAD_ROW0 = 44, 15


def banner(t: str) -> None:
    print()
    print("=" * 74)
    print(t)
    print("=" * 74)


# %%
# === Subject configuration ===
def load_subject(name: str) -> dict:
    """Read one subject's registry entry written by `scratch_cohort_registry.py`."""
    p = CONFIG_DIR / f"{name.lower()}.json"
    if not p.exists():
        raise FileNotFoundError(f"no registry for {name!r}; run scratch_cohort_registry.py")
    return json.loads(p.read_text(encoding="utf-8"))


def stream_ext(subject: str, role: str = "broadband") -> str | None:
    """Which file extension carries a given signal for this subject.

    ``broadband`` is ``.ns5`` for five subjects and ``.ns6`` for Fisk, so this
    must never be hardcoded. Returns None when the subject has no such stream.
    """
    return load_subject(subject)["streams"].get(role)


def implant_for(subject: str, date: str) -> dict | None:
    """The implant in place on a given date, by surgery order."""
    imps = load_subject(subject)["implants"]
    chosen = imps[0]
    for imp in imps:
        s = imp.get("surgery_date")
        if isinstance(s, str) and date >= s:
            chosen = imp
    return chosen


def implant_age_days(subject: str, date: str) -> float | None:
    """Days since this implant's own surgery, which is the longitudinal axis.

    Calendar date is the wrong axis for cross-subject comparison and actively
    misleading within Rocky, whose second implant starts at age zero in 2025.
    """
    imp = implant_for(subject, date)
    s = imp.get("surgery_date") if imp else None
    if not isinstance(s, str):
        return None
    return (pd.Timestamp(date) - pd.Timestamp(s)).days


# %%
# === Resilient workbook loading ===
def recover_truncated_xlsx(path: Path) -> io.BytesIO | None:
    """Rebuild a zip whose central directory is missing.

    `13966-8 SN 1025-001497.xlsm` is truncated: no end-of-central-directory
    record, so zipfile, openpyxl and Excel all refuse it. All 19 copies across
    8 volumes are byte-identical at 75,888 bytes, so it was truncated at source
    and there is no intact copy to fall back on.

    The central directory holds no content, only an index of members that each
    carry a complete local header. Walking those headers recovers the file
    losslessly, confirmed by the CRC stored in every member.
    """
    blob = path.read_bytes()
    members, pos = [], 0
    while (i := blob.find(b"PK\x03\x04", pos)) >= 0:
        if len(blob) < i + 30:
            break
        (_, _, flags, method, _, _, crc, csize, _, nlen, elen) = struct.unpack(
            "<IHHHHHIIIHH", blob[i:i + 30])
        name = blob[i + 30:i + 30 + nlen].decode("utf-8", "replace")
        start = i + 30 + nlen + elen
        if flags & 0x08:            # sizes live in a trailing descriptor
            pos = i + 4
            continue
        data = blob[start:start + csize]
        if len(data) < csize:
            break
        try:
            raw = zlib.decompress(data, -15) if method == 8 else data
        except zlib.error:
            pos = start + csize
            continue
        if zlib.crc32(raw) == crc:
            members.append((name, raw))
        pos = start + csize
    if not members:
        return None
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, raw in members:
            z.writestr(name, raw)
    buf.seek(0)
    buf.n_recovered = len(members)
    return buf


def open_workbook(path: Path):
    """Load a workbook, recovering it first if the archive is damaged."""
    try:
        return load_workbook(path, data_only=True), None
    except zipfile.BadZipFile:
        buf = recover_truncated_xlsx(path)
        if buf is None:
            return None, "unreadable: damaged archive, recovery failed"
        return (load_workbook(buf, data_only=True),
                f"RECOVERED from truncated archive ({buf.n_recovered} members, "
                f"all CRC-verified)")


# %%
# === Probe geometry ===
def parse_cmp(path: Path) -> pd.DataFrame:
    """Parse a Blackrock CMP into channel_id -> (col, row, bank, elec, label).

    Parameters
    ----------
    path : Path
        The `.cmp` mapfile shipped with the array.

    Returns
    -------
    pandas.DataFrame
        One row per electrode. ``channel_id = (bank - 'A') * 32 + elec`` is
        the NEV channel id; ``label`` is the manufacturer's ``elecN`` and is
        deliberately a different number (see utah_channel_mapping.md).
    """
    rows = []
    for ln in path.read_text(encoding="utf-8", errors="replace").splitlines():
        p = ln.split()
        if len(p) >= 4 and p[0].isdigit() and p[1].isdigit() and p[3].isdigit():
            col, row, bank, elec = int(p[0]), int(p[1]), p[2].upper(), int(p[3])
            rows.append(dict(
                col=col, row=row, bank=bank, elec=elec,
                label=p[4] if len(p) >= 5 else "",
                channel_id=(ord(bank) - ord("A")) * 32 + elec,
            ))
    return pd.DataFrame(rows)


def validate_cmp(cmp_df: pd.DataFrame) -> list[str]:
    """Internal-consistency checks that hold for any Utah-96, however wired.

    Returns human-readable problems; empty means the file is self-consistent.
    Deliberately says nothing about *which* four cells are vacant, because that
    is an individual array's build record rather than a property of the type.
    Use :func:`verify_against_padmap` for the geometry itself.
    """
    issues: list[str] = []
    n = len(cmp_df)
    if n != N_ELECTRODES:
        issues.append(f"{n} electrodes, expected {N_ELECTRODES}")

    pos = list(zip(cmp_df.col, cmp_df.row, strict=True))
    dupes = {p for p in pos if pos.count(p) > 1}
    if dupes:
        issues.append(f"duplicate grid positions: {sorted(dupes)}")

    out_of_range = cmp_df[(cmp_df.col >= GRID) | (cmp_df.row >= GRID)
                          | (cmp_df.col < 0) | (cmp_df.row < 0)]
    if len(out_of_range):
        issues.append(f"{len(out_of_range)} positions outside the {GRID}x{GRID} grid")

    if cmp_df.channel_id.duplicated().any():
        d = cmp_df.channel_id[cmp_df.channel_id.duplicated()].tolist()
        issues.append(f"duplicate channel_id: {sorted(set(d))}")
    if cmp_df.label.duplicated().any():
        d = cmp_df.label[cmp_df.label.duplicated()].tolist()
        issues.append(f"duplicate label: {sorted(set(d))}")
    bad_bank = set(cmp_df.bank) - set("ABCD")
    if bad_bank:
        issues.append(f"unexpected bank values: {sorted(bad_bank)}")
    bad_elec = cmp_df[(cmp_df.elec < 1) | (cmp_df.elec > 32)]
    if len(bad_elec):
        issues.append(f"{len(bad_elec)} elec values outside 1..32")
    return issues


def vacant_cells(cmp_df: pd.DataFrame) -> list[tuple[int, int]]:
    """The four grid cells this array does not populate."""
    occupied = set(zip(cmp_df.col, cmp_df.row, strict=True))
    return sorted({(c, r) for c in range(GRID) for r in range(GRID)} - occupied)


def read_pad_map(xlsm_path: Path) -> dict[tuple[int, int], str] | None:
    """Pad-side location grid from the factory workbook: (col, row) -> label.

    This is Blackrock's own record of where each electrode physically sits,
    printed as a 10x10 block viewed from the pad side. It is the authority for
    geometry -- the `.cmp` is a derived export of the same information.
    """
    wb, _ = open_workbook(xlsm_path)
    if wb is None or PAD_SHEET not in wb.sheetnames:
        return None
    ws = wb[PAD_SHEET]
    pad: dict[tuple[int, int], str] = {}
    for i in range(GRID):
        for j in range(GRID):
            v = ws.cell(row=PAD_ROW0 + i, column=PAD_COL0 + j).value
            if isinstance(v, (int, float)):
                pad[(j, GRID - 1 - i)] = f"elec{int(v)}"
    wb.close()
    return pad or None


def verify_against_padmap(cmp_df: pd.DataFrame,
                          pad: dict[tuple[int, int], str]) -> list[str]:
    """Compare a parsed CMP against the array's own pad-side location grid."""
    cmp_pos = {(int(r.col), int(r.row)): r.label for r in cmp_df.itertuples()}
    issues: list[str] = []
    only_pad = sorted(set(pad) - set(cmp_pos))
    only_cmp = sorted(set(cmp_pos) - set(pad))
    if only_pad:
        issues.append(f"populated in pad map but not in cmp: {only_pad}")
    if only_cmp:
        issues.append(f"populated in cmp but not in pad map: {only_cmp}")
    bad = [(k, pad[k], cmp_pos[k]) for k in set(pad) & set(cmp_pos)
           if pad[k] != cmp_pos[k]]
    for k, a, b in bad[:8]:
        issues.append(f"{k}: pad map says {a}, cmp says {b}")
    return issues


def load_probe_map(path: Path, xlsm: Path | None = None
                   ) -> tuple[pd.DataFrame, list[str]]:
    """Parse a CMP and check it, against the factory pad map where available.

    No repair step exists by design. A CMP that disagrees with a canonical
    layout is far more likely to be a rewired array than a corrupted file, and
    "fixing" it would move electrodes to cells that hold none.
    """
    df = parse_cmp(path)
    notes = [f"ISSUE: {i}" for i in validate_cmp(df)]
    vac = vacant_cells(df)
    if set(vac) != TYPICAL_VACANT:
        notes.append(f"NOTE: non-standard vacant cells {vac} "
                     f"(typical is {sorted(TYPICAL_VACANT)}) -- rewired array")
    if xlsm is not None and xlsm.exists():
        pad = read_pad_map(xlsm)
        if pad is None:
            notes.append("NOTE: no pad-side grid in the workbook")
        else:
            geo = verify_against_padmap(df, pad)
            notes += ([f"GEOMETRY MISMATCH: {g}" for g in geo] if geo
                      else [f"pad map agrees at {len(pad)}/{len(pad)} positions"])
    return df, notes


def find_cmp_files() -> list[tuple[Path, Path | None]]:
    """Every CMP available to the project, paired with its factory workbook."""
    out = []
    for p in sorted(PROBE_DIR.glob("*.cmp")):
        serial = re.search(r"(\d{4}-\d{6})", p.name)
        xl = None
        if serial:
            cands = list(PROBE_DIR.glob(f"*{serial.group(1)}.xlsm"))
            xl = cands[0] if cands else None
        out.append((p, xl))
    return out


# %%
# === Validation report ===
def run_validation() -> int:
    """Check every CMP against its own factory record, and report layouts."""
    files = find_cmp_files()
    banner(f"Validating {len(files)} CMP files against their factory pad maps")
    maps: dict[str, pd.DataFrame] = {}
    n_bad = 0
    for p, xl in files:
        df, notes = load_probe_map(p, xlsm=xl)
        serial = re.search(r"(\d{4}-\d{6})", p.name)
        key = serial.group(1) if serial else p.stem
        maps[key] = df
        bad = [n for n in notes if n.startswith(("ISSUE", "GEOMETRY"))]
        print(f"\n  {p.name:28s} {len(df):3d} electrodes   "
              f"{'DEFECT' if bad else 'OK'}"
              f"{'   (no workbook)' if xl is None else ''}")
        for n in notes:
            print(f"      {n}")
        if bad:
            n_bad += 1

    banner("Vacant cells per array")
    # Which four cells are empty is a build property. Arrays that lost shanks
    # during manufacture are rewired from surviving shanks elsewhere, so a
    # layout differing from its siblings is a fact about the array, not an
    # error in its mapfile.
    for k in sorted(maps):
        vac = vacant_cells(maps[k])
        kind = "typical" if set(vac) == TYPICAL_VACANT else "REWIRED"
        print(f"  {k:14s} {vac}   {kind}")

    banner("Summary")
    print(f"  {len(files) - n_bad} of {len(files)} CMP files consistent with "
          f"their own factory record")
    return 0


# %%
# === Regression against the session 4/5 pipeline ===
def run_regression() -> int:
    """Confirm the config-driven path reproduces the Rocky-specific one."""
    import sys

    sys.path.insert(0, str(REPO / "notebooks"))
    from scratch_rocky_resort import baseline_noise_uv, open_nev, read_electrode
    from scratch_rocky_spatial import CMP_BY_ARRAY
    from scratch_rocky_spatial import parse_cmp as old_parse_cmp

    banner("CMP parity: new parser vs the session 4/5 one")
    for arr, p in CMP_BY_ARRAY.items():
        if not Path(p).exists():
            print(f"  {arr:10s} SKIP (not staged): {p}")
            continue
        old = old_parse_cmp(Path(p)).sort_values("channel_id").reset_index(drop=True)
        new = parse_cmp(Path(p)).sort_values("channel_id").reset_index(drop=True)
        cols = ["col", "row", "bank", "elec", "label", "channel_id"]
        same = old[cols].equals(new[cols])
        print(f"  {arr:10s} {len(old)} electrodes   identical={same}")
        if not same:
            print(old[cols].compare(new[cols]).head(10).to_string())

    banner("Stream resolution from config")
    for s in ["Rocky", "Nigel", "Fisk", "Picasso"]:
        cfg = load_subject(s)
        print(f"  {s:9s} broadband={cfg['streams']['broadband']}  "
              f"lfp={cfg['streams']['lfp']}  snippets={cfg['streams']['snippets']}")

    banner("Implant assignment and age")
    for s, d in [("Rocky", "2018-04-26"), ("Rocky", "2024-03-29"),
                 ("Rocky", "2025-04-04"), ("Rocky", "2025-06-05"),
                 ("Nigel", "2023-03-17"), ("Fisk", "2024-06-19")]:
        imp = implant_for(s, d)
        age = implant_age_days(s, d)
        print(f"  {s:7s} {d}  -> {imp['implant']}  "
              f"arrays={list(imp['arrays'])}  "
              f"age={'unknown' if age is None else f'{age} d'}")

    banner("Loading one Rocky session through the config-driven path")
    idx_p = REPO / "data" / "derived" / "rocky" / "session_index.parquet"
    if not idx_p.exists():
        print("  SKIP: Rocky session index not present")
        return 0
    idx = pd.read_parquet(idx_p)
    row = idx[(idx.date == "2018-04-26") & (idx.array == "Anterior")
              & (idx.kind == "OFS")]
    if not len(row):
        print("  SKIP: reference session not in index")
        return 0
    p = Path(row.iloc[0]["path"])
    raw, meta, cbe = open_nev(p)
    noises = []
    for elec in sorted(cbe):
        e = read_electrode(raw, meta, cbe[elec])
        if e is not None and len(e["t"]) >= 50:
            noises.append(baseline_noise_uv(e["wf"], meta["nbefore"]))
    print(f"  {p.name}")
    print(f"    electrodes {len(noises)}   median noise {np.median(noises):.2f} uV")
    print("    session 5 reported: 96 electrodes, median noise 9.82 uV")
    ok = len(noises) == 96 and abs(np.median(noises) - 9.82) < 0.01
    print(f"    REGRESSION {'PASS' if ok else 'FAIL'}")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--validate", action="store_true")
    ap.add_argument("--regress", action="store_true")
    args = ap.parse_args()
    if args.validate:
        run_validation()
    if args.regress:
        run_regression()
    if not (args.validate or args.regress):
        run_validation()
        run_regression()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
